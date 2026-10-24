"""把 食材總表.xlsx 轉成網頁用的 data/ingredients.json。

用法：python3 convert.py <食材總表.xlsx 路徑> <輸出 json 路徑>

只輸出 ID、類別、品項、每克價格（不含廠商、進貨價、備註）。
若有「別名」欄位（表頭寫「別名」），以頓號/逗號分隔的別名也會一起輸出。
"""
import datetime
import json
import os
import re
import sys

import openpyxl

src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
ws = wb["工作表1"] if "工作表1" in wb.sheetnames else wb.worksheets[0]

rows = ws.iter_rows(values_only=True)
header = [str(h or "").strip() for h in next(rows)]


def col(name, default):
    for i, h in enumerate(header):
        if h.startswith(name):
            return i
    return default


c_id, c_cat, c_name, c_ppg = col("ID", 0), col("類別", 1), col("項目", 2), col("每克價格", 8)
c_alias = col("別名", None)

items = []
for r in rows:
    name = r[c_name] if len(r) > c_name else None
    if not name or not str(name).strip():
        continue
    rid = str(r[c_id] or "").strip()
    if rid and not (rid.startswith("ID") or rid.startswith("PREP")):
        continue
    try:
        ppg = float(r[c_ppg]) if r[c_ppg] not in (None, "") else 0.0
    except (TypeError, ValueError):
        ppg = 0.0
    it = {"id": rid, "cat": str(r[c_cat] or "").strip(), "name": str(name).strip(), "ppg": round(ppg, 6)}
    if c_alias is not None and len(r) > c_alias and r[c_alias]:
        it["alias"] = [a.strip() for a in re.split(r"[、,，/]", str(r[c_alias])) if a.strip()]
    items.append(it)

TPE = datetime.timezone(datetime.timedelta(hours=8))
updated = datetime.datetime.fromtimestamp(os.path.getmtime(src), tz=TPE).strftime("%Y-%m-%d %H:%M")
os.makedirs(os.path.dirname(os.path.abspath(dst)), exist_ok=True)
with open(dst, "w", encoding="utf-8") as f:
    json.dump({"updated": updated, "items": items}, f, ensure_ascii=False, separators=(",", ":"))
print(f"{len(items)} items, source modified {updated}")
